"""
app.py - Resume Keyword Analyzer (Modern Streamlit Web Application)
===================================================================
A modern, responsive web application built with Streamlit, Python, and NLP-based
text processing. Enables candidates to upload resumes (PDF, DOCX, TXT), compare
them against job descriptions, review match scores, inspect skill gaps, and explore
personalized upskilling roadmaps.
"""

import io
import os
from typing import Tuple, Dict, Any, List

import streamlit as st
import pypdf
import docx

from skill_extractor import (
    extract_skills,
    calculate_match,
    get_all_skills_by_category,
    SKILL_DICTIONARY,
)
from roadmap import generate_roadmap, RoadmapPlan


# Paths to bundled sample files
SAMPLE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples")
SAMPLE_RESUME_PATH = os.path.join(SAMPLE_DIR, "sample_resume.txt")
SAMPLE_JOB_PATH = os.path.join(SAMPLE_DIR, "sample_job_description.txt")


def load_bundled_sample(filepath: str) -> str:
    """Helper to load bundled sample text."""
    if os.path.isfile(filepath):
        try:
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        except Exception:
            return ""
    return ""


def extract_text_from_upload(uploaded_file) -> Tuple[bool, str]:
    """
    Safely extracts plain text from uploaded PDF, DOCX, or TXT files.
    Returns:
        (True, extracted_text) on success
        (False, error_message) on failure
    """
    if uploaded_file is None:
        return False, "No file uploaded."

    filename = uploaded_file.name.lower()

    try:
        if filename.endswith(".pdf"):
            pdf_bytes = io.BytesIO(uploaded_file.read())
            reader = pypdf.PdfReader(pdf_bytes)

            if reader.is_encrypted:
                try:
                    # Attempt empty password decryption
                    reader.decrypt("")
                except Exception:
                    return False, "This PDF is encrypted and protected by a password."

            text_chunks: List[str] = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_chunks.append(page_text)

            full_text = "\n".join(text_chunks).strip()
            if not full_text:
                return (
                    False,
                    "No readable text could be extracted from this PDF. It may be a scanned image-only document.",
                )
            return True, full_text

        elif filename.endswith(".docx"):
            docx_bytes = io.BytesIO(uploaded_file.read())
            doc = docx.Document(docx_bytes)

            text_chunks: List[str] = []
            for p in doc.paragraphs:
                if p.text.strip():
                    text_chunks.append(p.text)

            # Also extract text from any tables in the DOCX
            for table in doc.tables:
                for row in table.rows:
                    row_content = [
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    ]
                    if row_content:
                        text_chunks.append(" | ".join(row_content))

            full_text = "\n".join(text_chunks).strip()
            if not full_text:
                return False, "The DOCX file did not contain any readable text."
            return True, full_text

        elif filename.endswith(".txt"):
            raw_bytes = uploaded_file.read()
            # Attempt UTF-8 decoding, fallback to latin-1
            try:
                full_text = raw_bytes.decode("utf-8")
            except UnicodeDecodeError:
                full_text = raw_bytes.decode("latin-1", errors="replace")

            full_text = full_text.strip()
            if not full_text:
                return False, "The TXT file is empty."
            return True, full_text

        else:
            return (
                False,
                f"Unsupported format '{uploaded_file.name}'. Please upload PDF, DOCX, or TXT.",
            )

    except Exception as exc:
        return False, f"Failed to extract text from '{uploaded_file.name}': {str(exc)}"


# =====================================================================
# Main Streamlit UI Setup
# =====================================================================

def run_app():
    """Main rendering function for Streamlit web app."""
    st.set_page_config(
        page_title="Resume Keyword Analyzer",
        page_icon="📄",
        layout="wide",
        initial_sidebar_state="collapsed",
    )

    # Custom CSS for Modern Blue-and-White Theme & Clean Component Cards
    custom_css = """
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
        
        html, body, [class*="css"] {
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif;
        }
        
        .stApp {
            background: linear-gradient(180deg, #F8FAFC 0%, #EFF6FF 100%);
            color: #0F172A;
        }

        .block-container {
            padding-top: 2rem;
            padding-bottom: 4rem;
            max-width: 1200px;
        }

        /* Modern Blue Hero Header */
        .hero-container {
            background: linear-gradient(135deg, #1E3A8A 0%, #2563EB 50%, #3B82F6 100%);
            border-radius: 18px;
            padding: 38px 30px;
            color: #FFFFFF;
            box-shadow: 0 10px 25px -5px rgba(37, 99, 235, 0.25);
            margin-bottom: 28px;
            text-align: center;
            border: 1px solid rgba(255, 255, 255, 0.15);
        }
        .hero-title {
            font-size: 2.5rem;
            font-weight: 800;
            margin-bottom: 8px;
            letter-spacing: -0.025em;
            color: #FFFFFF;
        }
        .hero-subtitle {
            font-size: 1.15rem;
            color: #DBEAFE;
            font-weight: 400;
            max-width: 650px;
            margin: 0 auto 18px auto;
            line-height: 1.5;
        }
        .hero-tags {
            display: flex;
            justify-content: center;
            gap: 12px;
            flex-wrap: wrap;
        }
        .hero-tag {
            background: rgba(255, 255, 255, 0.16);
            backdrop-filter: blur(8px);
            padding: 6px 14px;
            border-radius: 20px;
            font-size: 0.85rem;
            font-weight: 500;
            border: 1px solid rgba(255, 255, 255, 0.2);
        }

        /* Card Containers */
        .custom-card {
            background: #FFFFFF;
            border-radius: 16px;
            border: 1px solid #E2E8F0;
            padding: 24px;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.04), 0 2px 4px -1px rgba(0, 0, 0, 0.02);
            margin-bottom: 24px;
            transition: transform 0.2s ease, box-shadow 0.2s ease;
        }
        .custom-card:hover {
            box-shadow: 0 10px 15px -3px rgba(0, 0, 0, 0.07);
        }

        .card-header {
            display: flex;
            align-items: center;
            gap: 10px;
            font-size: 1.25rem;
            font-weight: 700;
            color: #1E3A8A;
            margin-bottom: 16px;
            border-bottom: 2px solid #F1F5F9;
            padding-bottom: 10px;
        }

        /* Primary CTA Button */
        div.stButton > button:first-child {
            background: linear-gradient(135deg, #2563EB 0%, #1D4ED8 100%);
            color: #FFFFFF;
            font-size: 1.1rem;
            font-weight: 700;
            padding: 14px 28px;
            border-radius: 12px;
            border: none;
            box-shadow: 0 4px 14px 0 rgba(37, 99, 235, 0.35);
            transition: all 0.2s ease;
            width: 100%;
            margin-top: 10px;
            margin-bottom: 20px;
        }
        div.stButton > button:first-child:hover {
            background: linear-gradient(135deg, #1D4ED8 0%, #1E40AF 100%);
            box-shadow: 0 6px 20px 0 rgba(37, 99, 235, 0.5);
            transform: translateY(-2px);
        }

        /* Metric Score Highlight Box */
        .score-banner {
            background: #FFFFFF;
            border-radius: 18px;
            border: 1px solid #E2E8F0;
            padding: 28px;
            box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.06);
            margin-bottom: 28px;
        }
        .score-number {
            font-size: 3.5rem;
            font-weight: 800;
            line-height: 1;
            margin-bottom: 4px;
        }
        .score-label {
            font-size: 1rem;
            font-weight: 600;
            color: #64748B;
            text-transform: uppercase;
            letter-spacing: 0.05em;
        }
        .score-badge {
            display: inline-block;
            padding: 6px 16px;
            border-radius: 9999px;
            font-size: 0.95rem;
            font-weight: 700;
            margin-top: 10px;
        }

        /* Badges / Pill Styles */
        .badge-matched {
            display: inline-flex;
            align-items: center;
            gap: 6px;
            background-color: #ECFDF5;
            color: #065F46;
            border: 1px solid #A7F3D0;
            padding: 5px 12px;
            border-radius: 9999px;
            font-size: 0.88rem;
            font-weight: 600;
            margin: 4px;
        }
        .badge-missing {
            display: inline-flex;
            align-items: center;
            gap: 6px;
            background-color: #FEF2F2;
            color: #991B1B;
            border: 1px solid #FECACA;
            padding: 5px 12px;
            border-radius: 9999px;
            font-size: 0.88rem;
            font-weight: 600;
            margin: 4px;
        }
        .badge-bonus {
            display: inline-flex;
            align-items: center;
            gap: 6px;
            background-color: #EFF6FF;
            color: #1E40AF;
            border: 1px solid #BFDBFE;
            padding: 5px 12px;
            border-radius: 9999px;
            font-size: 0.88rem;
            font-weight: 600;
            margin: 4px;
        }

        /* Roadmap Phase Cards */
        .phase-card {
            background: #F8FAFC;
            border-radius: 14px;
            border: 1px solid #E2E8F0;
            padding: 20px;
            margin-bottom: 16px;
        }
        .phase-title {
            font-size: 1.15rem;
            font-weight: 700;
            color: #1E3A8A;
            margin-bottom: 12px;
            display: flex;
            align-items: center;
            gap: 8px;
        }
        .skill-card {
            background: #FFFFFF;
            border-radius: 10px;
            border: 1px solid #E2E8F0;
            padding: 14px 18px;
            margin-bottom: 10px;
        }
        .skill-card-top {
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 6px;
        }
        .skill-card-name {
            font-size: 1.05rem;
            font-weight: 700;
            color: #0F172A;
        }
        .skill-card-priority {
            font-size: 0.78rem;
            font-weight: 700;
            padding: 3px 10px;
            border-radius: 9999px;
        }
        .priority-high {
            background: #FEE2E2;
            color: #991B1B;
        }
        .priority-medium-high {
            background: #FEF3C7;
            color: #92400E;
        }
        .priority-medium {
            background: #E0E7FF;
            color: #3730A3;
        }

        /* Disclaimer Box */
        .disclaimer-box {
            background: #EFF6FF;
            border-left: 5px solid #3B82F6;
            border-radius: 8px;
            padding: 16px 20px;
            font-size: 0.9rem;
            color: #1E3A8A;
            line-height: 1.5;
            margin-top: 24px;
        }
    </style>
    """
    st.markdown(custom_css, unsafe_allow_html=True)

    # -------------------------------------------------------------
    # Dashboard Hero Header
    # -------------------------------------------------------------
    st.markdown(
        """
        <div class="hero-container">
            <div class="hero-title">📄 Resume Keyword Analyzer</div>
            <div class="hero-subtitle">Optimize your resume for your dream job. Match required skills, identify critical gaps, and accelerate your career with a tailored learning roadmap.</div>
            <div class="hero-tags">
                <span class="hero-tag">⚡ NLP Text Processing</span>
                <span class="hero-tag">🎯 33+ Technical & Soft Skills</span>
                <span class="hero-tag">📊 Match Score & Gap Audit</span>
                <span class="hero-tag">🗺️ Personalized Learning Roadmap</span>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    # -------------------------------------------------------------
    # Upload Cards: Resume & Job Description
    # -------------------------------------------------------------
    col_resume, col_job = st.columns([1, 1], gap="large")

    with col_resume:
        st.markdown(
            """
            <div class="custom-card">
                <div class="card-header">
                    <span>📁</span> Step 1: Upload Your Resume
                </div>
                <p style="color: #64748B; font-size: 0.95rem; margin-bottom: 12px;">
                    Upload your resume in <strong>PDF</strong>, <strong>DOCX</strong>, or <strong>TXT</strong> format.
                </p>
            </div>
            """,
            unsafe_allow_html=True,
        )

        resume_file = st.file_uploader(
            "Choose your resume file",
            type=["pdf", "docx", "txt"],
            help="Supports PDF, DOCX, or TXT files.",
            key="resume_uploader",
        )

        use_sample_resume = st.checkbox(
            "💡 Or use built-in sample candidate resume (Alex Morgan, Full-Stack)",
            value=False,
            key="sample_resume_checkbox",
        )

        if resume_file is not None:
            st.success(f"✔ Attached: **{resume_file.name}** ({round(resume_file.size / 1024, 1)} KB)")
        elif use_sample_resume:
            st.info("✔ Loaded built-in sample resume: Alex Morgan (3+ YOE Software Engineer)")

    with col_job:
        st.markdown(
            """
            <div class="custom-card">
                <div class="card-header">
                    <span>🎯</span> Step 2: Target Job Description
                </div>
                <p style="color: #64748B; font-size: 0.95rem; margin-bottom: 12px;">
                    Paste the job requirements or upload a target job posting (.txt).
                </p>
            </div>
            """,
            unsafe_allow_html=True,
        )

        tab_paste, tab_upload = st.tabs(["✍️ Paste Job Description", "📂 Upload TXT File"])

        job_text_input = ""
        with tab_paste:
            use_sample_job = st.checkbox(
                "💡 Or load built-in sample job description (CloudScale Innovations)",
                value=False,
                key="sample_job_checkbox",
            )
            default_job_val = load_bundled_sample(SAMPLE_JOB_PATH) if use_sample_job else ""

            job_text_input = st.text_area(
                "Job Description Text",
                value=default_job_val,
                height=210,
                placeholder="Paste responsibilities, required qualifications, and tech stack here...",
                key="job_text_area",
                label_visibility="collapsed",
            )

        job_file_upload = None
        with tab_upload:
            job_file_upload = st.file_uploader(
                "Choose a job description file (.txt)",
                type=["txt"],
                key="job_file_uploader",
            )

    # -------------------------------------------------------------
    # Prominent CTA Button
    # -------------------------------------------------------------
    st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True)
    analyze_clicked = st.button("🚀 Analyze Resume & Generate Roadmap", use_container_width=True)

    # -------------------------------------------------------------
    # Analysis & Results Rendering (Directly on Webpage, No Terminal Print)
    # -------------------------------------------------------------
    if analyze_clicked:
        resume_content = ""
        if resume_file is not None:
            with st.spinner("Extracting text from uploaded resume..."):
                ok_res, extracted_res = extract_text_from_upload(resume_file)
                if not ok_res:
                    st.error(f"❌ Resume Error: {extracted_res}")
                    return
                resume_content = extracted_res
        elif use_sample_resume:
            resume_content = load_bundled_sample(SAMPLE_RESUME_PATH)
            if not resume_content:
                st.error("❌ Could not load the sample resume file.")
                return
        else:
            st.warning("⚠️ **Validation Notice**: Please upload a resume file (PDF, DOCX, or TXT) or check 'Use sample resume'.")
            return

        job_content = ""
        if job_file_upload is not None:
            with st.spinner("Reading job description file..."):
                ok_job, extracted_job = extract_text_from_upload(job_file_upload)
                if not ok_job:
                    st.error(f"❌ Job Description Error: {extracted_job}")
                    return
                job_content = extracted_job
        elif job_text_input.strip():
            job_content = job_text_input.strip()
        else:
            st.warning("⚠️ **Validation Notice**: Please paste the job description text or upload a job description file.")
            return

        with st.spinner("Analyzing keywords, matching competencies, and building your personalized roadmap..."):
            resume_skills_dict = extract_skills(resume_content)
            job_skills_dict = extract_skills(job_content)

            resume_skill_set = set(resume_skills_dict.keys())
            job_skill_set = set(job_skills_dict.keys())

            match_stats = calculate_match(resume_skill_set, job_skill_set)
            roadmap: RoadmapPlan = generate_roadmap(
                missing_skills=match_stats["missing_skills"],
                matched_skills=match_stats["matched_skills"],
                candidate_skills=list(resume_skill_set),
                match_percentage=match_stats["match_percentage"],
            )

        st.markdown("<hr style='margin: 35px 0 25px 0; border: none; height: 1px; background-color: #E2E8F0;'>", unsafe_allow_html=True)
        st.markdown("<h2 style='color: #1E3A8A; font-weight: 800; margin-bottom: 20px;'>📊 Analysis & Match Results</h2>", unsafe_allow_html=True)

        if not match_stats["has_job_skills"]:
            st.warning(
                """
                ⚠️ **Notice: No Recognized Technical Skills Detected in Job Description**
                
                The provided job description did not match any of the 33+ keywords in our skill dictionary.
                Please ensure the job description contains explicit technical requirements (e.g., Python, SQL, React, AWS, Docker),
                or add custom domain terms to `skill_extractor.py`.
                """
            )
            return

        score = match_stats["match_percentage"]
        if score >= 80.0:
            score_color = "#16A34A"
            badge_bg = "#DCFCE7"
            badge_color = "#166534"
            badge_text = "Exceptional Match"
        elif score >= 65.0:
            score_color = "#2563EB"
            badge_bg = "#DBEAFE"
            badge_color = "#1E40AF"
            badge_text = "Strong Candidate"
        elif score >= 50.0:
            score_color = "#D97706"
            badge_bg = "#FEF3C7"
            badge_color = "#92400E"
            badge_text = "Moderate Match"
        else:
            score_color = "#DC2626"
            badge_bg = "#FEE2E2"
            badge_color = "#991B1B"
            badge_text = "Skills Gap Identified"

        st.markdown(
            f"""
            <div class="score-banner">
                <div style="display: flex; flex-wrap: wrap; align-items: center; justify-content: space-between; gap: 20px;">
                    <div style="flex: 1; min-width: 250px;">
                        <span class="score-label">Overall Match Score</span>
                        <div class="score-number" style="color: {score_color};">{score:.1f}%</div>
                        <div class="score-badge" style="background-color: {badge_bg}; color: {badge_color};">
                            {badge_text}
                        </div>
                        <p style="color: #475569; font-size: 0.95rem; margin-top: 10px; line-height: 1.4;">
                            {roadmap.get_summary_level()}
                        </p>
                    </div>
                    <div style="flex: 2; min-width: 320px;">
                        <div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(130px, 1fr)); gap: 14px;">
                            <div style="background: #F8FAFC; border: 1px solid #E2E8F0; border-radius: 12px; padding: 14px; text-align: center;">
                                <div style="font-size: 1.8rem; font-weight: 800; color: #1E3A8A;">{match_stats['total_job_skills']}</div>
                                <div style="font-size: 0.8rem; color: #64748B; font-weight: 600;">Job Skills Found</div>
                            </div>
                            <div style="background: #F8FAFC; border: 1px solid #E2E8F0; border-radius: 12px; padding: 14px; text-align: center;">
                                <div style="font-size: 1.8rem; font-weight: 800; color: #1E3A8A;">{match_stats['total_resume_skills']}</div>
                                <div style="font-size: 0.8rem; color: #64748B; font-weight: 600;">Resume Skills</div>
                            </div>
                            <div style="background: #ECFDF5; border: 1px solid #A7F3D0; border-radius: 12px; padding: 14px; text-align: center;">
                                <div style="font-size: 1.8rem; font-weight: 800; color: #065F46;">{match_stats['match_count']}</div>
                                <div style="font-size: 0.8rem; color: #047857; font-weight: 600;">Matched Skills</div>
                            </div>
                            <div style="background: #FEF2F2; border: 1px solid #FECACA; border-radius: 12px; padding: 14px; text-align: center;">
                                <div style="font-size: 1.8rem; font-weight: 800; color: #991B1B;">{match_stats['missing_count']}</div>
                                <div style="font-size: 0.8rem; color: #B91C1C; font-weight: 600;">Missing Gaps</div>
                            </div>
                            <div style="background: #EFF6FF; border: 1px solid #BFDBFE; border-radius: 12px; padding: 14px; text-align: center;">
                                <div style="font-size: 1.8rem; font-weight: 800; color: #1E40AF;">{len(match_stats['bonus_skills'])}</div>
                                <div style="font-size: 0.8rem; color: #1D4ED8; font-weight: 600;">Bonus Skills</div>
                            </div>
                        </div>
                    </div>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        st.progress(score / 100.0)

        # Matched and Missing Breakdown Cards
        col_matched, col_missing = st.columns([1, 1], gap="large")

        with col_matched:
            st.markdown(
                f"""
                <div class="custom-card">
                    <div class="card-header" style="color: #065F46;">
                        <span>✔</span> Matched Skills ({match_stats['match_count']})
                    </div>
                    <p style="color: #64748B; font-size: 0.9rem; margin-bottom: 14px;">
                        Skills in your resume that directly fulfill the job posting requirements:
                    </p>
                """,
                unsafe_allow_html=True,
            )

            if match_stats["matched_skills"]:
                cat_matched: Dict[str, List[str]] = {}
                for s in match_stats["matched_skills"]:
                    c = SKILL_DICTIONARY[s].category
                    cat_matched.setdefault(c, []).append(s)

                for cat, s_list in sorted(cat_matched.items()):
                    st.markdown(f"<div style='font-weight: 700; font-size: 0.85rem; color: #475569; margin: 10px 0 4px 4px;'>{cat}</div>", unsafe_allow_html=True)
                    badges_html = " ".join([f"<span class='badge-matched'>✓ {s}</span>" for s in sorted(s_list)])
                    st.markdown(f"<div>{badges_html}</div>", unsafe_allow_html=True)
            else:
                st.info("No matching skills detected between your resume and this job description.")

            st.markdown("</div>", unsafe_allow_html=True)

        with col_missing:
            st.markdown(
                f"""
                <div class="custom-card">
                    <div class="card-header" style="color: #991B1B;">
                        <span>✖</span> Missing Required Skills ({match_stats['missing_count']})
                    </div>
                    <p style="color: #64748B; font-size: 0.9rem; margin-bottom: 14px;">
                        Required by the job description but not detected on your resume:
                    </p>
                """,
                unsafe_allow_html=True,
            )

            if match_stats["missing_skills"]:
                cat_missing: Dict[str, List[str]] = {}
                for s in match_stats["missing_skills"]:
                    c = SKILL_DICTIONARY[s].category
                    cat_missing.setdefault(c, []).append(s)

                for cat, s_list in sorted(cat_missing.items()):
                    st.markdown(f"<div style='font-weight: 700; font-size: 0.85rem; color: #475569; margin: 10px 0 4px 4px;'>{cat}</div>", unsafe_allow_html=True)
                    badges_html = " ".join([f"<span class='badge-missing'>! {s}</span>" for s in sorted(s_list)])
                    st.markdown(f"<div>{badges_html}</div>", unsafe_allow_html=True)
            else:
                st.success("🎉 Phenomenal! 100% of required technical skills are present on your resume.")

            st.markdown("</div>", unsafe_allow_html=True)

        if match_stats["bonus_skills"]:
            st.markdown(
                f"""
                <div class="custom-card">
                    <div class="card-header" style="color: #1E40AF;">
                        <span>★</span> Additional Candidate Strengths ({len(match_stats['bonus_skills'])})
                    </div>
                    <p style="color: #64748B; font-size: 0.9rem; margin-bottom: 10px;">
                        Bonus skills present on your resume that were not explicitly listed in the job description:
                    </p>
                """,
                unsafe_allow_html=True,
            )
            bonus_badges = " ".join([f"<span class='badge-bonus'>★ {s}</span>" for s in match_stats["bonus_skills"]])
            st.markdown(f"<div>{bonus_badges}</div>", unsafe_allow_html=True)
            st.markdown("</div>", unsafe_allow_html=True)

        # Personalized Learning Roadmap
        if match_stats["missing_skills"]:
            st.markdown("<h3 style='color: #1E3A8A; font-weight: 800; margin-top: 30px;'>🗺️ Personalized Learning Roadmap</h3>", unsafe_allow_html=True)
            st.markdown(
                f"""
                <p style="color: #475569; font-size: 1rem; margin-bottom: 16px;">
                    Estimated dedicated study duration: <strong>~{roadmap.total_estimated_weeks} weeks</strong>.
                    Missing skills are sequenced below in order of foundation, frameworks, and cloud specializations:
                </p>
                """,
                unsafe_allow_html=True,
            )

            for phase_name, skill_cards in roadmap.phases.items():
                st.markdown(
                    f"""
                    <div class="phase-card">
                        <div class="phase-title">📌 {phase_name}</div>
                    """,
                    unsafe_allow_html=True,
                )

                for card in skill_cards:
                    priority_class = (
                        "priority-high"
                        if card["tier"] == 1
                        else "priority-medium-high"
                        if card["tier"] == 2
                        else "priority-medium"
                    )

                    prereq_alert = ""
                    if card["unmet_prerequisites"]:
                        prereq_names = ", ".join(card["unmet_prerequisites"])
                        prereq_alert = f"<div style='color: #B45309; font-size: 0.84rem; margin-top: 4px;'>⚠️ <em>Recommended prerequisite to learn first: {prereq_names}</em></div>"

                    topics_html = ""
                    if card["key_topics"]:
                        topics_html = f"<div style='color: #334155; font-size: 0.88rem; margin-top: 6px;'><strong>Core Concepts:</strong> {', '.join(card['key_topics'])}</div>"

                    project_html = ""
                    if card["project_idea"]:
                        project_html = f"<div style='color: #1E40AF; font-size: 0.88rem; margin-top: 4px;'><strong>🛠️ Mini-Project Idea:</strong> {card['project_idea']}</div>"

                    st.markdown(
                        f"""
                        <div class="skill-card">
                            <div class="skill-card-top">
                                <span class="skill-card-name">{card['name']} <span style="font-size: 0.85rem; font-weight: 500; color: #64748B;">({card['category']})</span></span>
                                <span class="skill-card-priority {priority_class}">{card['priority']}</span>
                            </div>
                            <div style="color: #475569; font-size: 0.9rem; line-height: 1.4;">
                                {card['description']}
                            </div>
                            {prereq_alert}
                            {topics_html}
                            {project_html}
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )

                st.markdown("</div>", unsafe_allow_html=True)

        # Strategic Resume Suggestions
        st.markdown("<h3 style='color: #1E3A8A; font-weight: 800; margin-top: 30px;'>💡 Strategic Resume Improvement Suggestions</h3>", unsafe_allow_html=True)
        st.markdown(
            """
            <div class="custom-card">
                <ol style="margin-left: 20px; line-height: 1.7; color: #334155; font-size: 0.95rem;">
                    <li><strong>Highlight Transferable & Adjacent Experience:</strong> If you are missing a framework like FastAPI but have extensive experience in Flask, mention: <em>"Proficient in Python web services (Flask); rapidly adopting FastAPI for asynchronous schemas."</em></li>
                    <li><strong>Demonstrate Missing Skills via GitHub:</strong> Recruiters value tangible proof. Build 1-2 focused repository projects showcasing your missing skills with clean documentation, tests, and CI/CD pipelines.</li>
                    <li><strong>Use the STAR Method with Quantifiable Impact:</strong> Structure experience bullets with <em>Situation, Task, Action, Result</em> and quantify outcomes (e.g. <em>"Optimized PostgreSQL queries using SQL indexing, cutting latency by 35%"</em>).</li>
                    <li><strong>Contextual Keyword Placement:</strong> Don't just list skills in a bottom keyword dump. Ensure critical job keywords appear both in your <em>Technical Skills summary</em> and woven directly into your <em>Experience bullet points</em>.</li>
                </ol>
            </div>
            """,
            unsafe_allow_html=True,
        )

        # ATS Disclaimer
        st.markdown(
            """
            <div class="disclaimer-box">
                <strong>⚠️ Transparency Note: Keyword Estimate vs. Official ATS Scoring</strong><br>
                This score is an automated keyword-based heuristic estimate based on detected skills. It is <strong>not</strong> an official Applicant Tracking System (ATS) score. Commercial ATS platforms (e.g., Workday, Greenhouse, Taleo) and human recruiters also assess career trajectory, years of experience, job title relevancy, education, and contextual impact. Use this analysis as a tactical preparation checklist to ensure your resume accurately conveys your technical strengths.
            </div>
            """,
            unsafe_allow_html=True,
        )

        # Download Report Feature
        def generate_download_report() -> str:
            report_lines = [
                "============================================================================",
                "                   RESUME KEYWORD ANALYSIS REPORT                   ",
                "============================================================================",
                f"Overall Match Score:       {score:.1f}%",
                f"Candidate Assessment:      {roadmap.get_summary_level()}",
                f"Total Job Skills:          {match_stats['total_job_skills']}",
                f"Total Resume Skills:       {match_stats['total_resume_skills']}",
                f"Matched Skills:            {match_stats['match_count']}",
                f"Missing Skills:            {match_stats['missing_count']}",
                f"Bonus Resume Skills:       {len(match_stats['bonus_skills'])}",
                "----------------------------------------------------------------------------",
                "MATCHED SKILLS:",
                ", ".join(match_stats["matched_skills"]) if match_stats["matched_skills"] else "None",
                "",
                "MISSING SKILLS:",
                ", ".join(match_stats["missing_skills"]) if match_stats["missing_skills"] else "None",
                "",
                "BONUS SKILLS:",
                ", ".join(match_stats["bonus_skills"]) if match_stats["bonus_skills"] else "None",
                "============================================================================",
                "LEARNING ROADMAP:",
                f"Estimated Study Timeline: ~{roadmap.total_estimated_weeks} weeks",
            ]
            for p_name, cards in roadmap.phases.items():
                report_lines.append(f"\n[{p_name.upper()}]")
                for c in cards:
                    report_lines.append(f"  • {c['name']} ({c['priority']}) - {c['estimated_weeks']} week(s)")
                    report_lines.append(f"    Project: {c['project_idea']}")
            report_lines.append("============================================================================")
            return "\n".join(report_lines)

        st.download_button(
            label="📥 Download Complete Analysis Report (.txt)",
            data=generate_download_report(),
            file_name="resume_keyword_analysis_report.txt",
            mime="text/plain",
            use_container_width=True,
        )


if __name__ == "__main__":
    run_app()
